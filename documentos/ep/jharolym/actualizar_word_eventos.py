"""Incorpora la vista de eventos al Word local sin sobrescribir el original abierto.

Requiere python-docx y el PNG generado desde eventos-crm-erp-ep.svg.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "GE703_EP_G7_JHAROLYM_ARQUITECTURA_INTEGRADA.docx"
DEST = ROOT / "GE703_EP_G7_JHAROLYM_ARQUITECTURA_INTEGRADA_CON_EVENTOS.docx"
IMAGE = ROOT / "arquitectura/diagramas/eventos-crm-erp-ep.png"


def paragraph_before(anchor, text, style="Normal", page_break=False, size=11):
    paragraph = anchor.insert_paragraph_before(text, style)
    fmt = paragraph.paragraph_format
    fmt.page_break_before = page_break
    fmt.space_after = Pt(7)
    fmt.line_spacing = 1.08
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(size if style == "Normal" else 12)
        run.font.bold = style != "Normal"
        run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def events_table(document, anchor):
    records = [
        ["E01", "Cotización aceptada\nCliente / Ventas C03", "Vendedor en CRM C01 puede solicitar evaluación por I03; C04 contrasta versión y aceptación por I04.", "ACE-EP-001 en v1. Antecedente sintético; no lo ejecuta el simulador."],
        ["E02", "Solicitud recibida\nCRM C01 por I03", "C04 comprueba permisos e idempotencia y coordina las validaciones.", "RECIBIDA / RECIBIDA en la bitácora de los tres casos."],
        ["E03", "Evaluación validada\nControl C04", "CRM / Ventas recibe ACEPTADA con código VALIDADA y comunica el resultado.", "DECISION / VALIDADA en EP-N01; order_id = null."],
        ["E04", "Disponibilidad parcial\nC05 informa; C04 decide", "CRM / Ventas recibe PROMESA_CONDICIONADA y revisa condiciones con el cliente.", "DECISION / DISPONIBILIDAD_PARCIAL en EP-S01; brecha de 2 unidades."],
        ["E05", "Descuento sin aprobación\nC06 se consulta; C04 decide", "CRM / Ventas recibe RECHAZADA y escala al dueño. El rechazo precede a la consulta de stock.", "DECISION / DESCUENTO_SIN_APROBACION en EP-E01; cotización v2."],
        ["E06", "Necesidad de abastecer\nArnés tras respuesta parcial", "Dueño / abastecimiento evalúa producto, ubicación y faltante. Conexión conceptual O2C–P2P.", "REP-EP-001 en EP-S01, fuera del contrato 1.0; automatic_action = false."],
        ["E07", "Solicitud final repetida\nCRM C01 por I03", "C04 devuelve la respuesta original para la misma clave y contenido. No vuelve a evaluar.", "REINTENTO y código original en los tres casos; una sola decisión persistida."],
    ]
    table = document.add_table(rows=1, cols=4)
    table.autofit = False
    widths = [1.0, 4.15, 6.05, 6.53]
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    headers = ["ID", "Evento y origen", "Destinatario y acción", "Registro y evidencia"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for record in records:
        for cell, text in zip(table.add_row().cells, record):
            cell.text = text
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        for name, value in (("val", "single"), ("sz", "5"), ("color", "D9D9D9")):
            element.set(qn(f"w:{name}"), value)
        borders.append(element)
    table._tbl.tblPr.append(borders)
    for i, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for j, cell in enumerate(row.cells):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            fill = OxmlElement("w:shd")
            fill.set(qn("w:fill"), "1F4D64" if i == 0 else ("F0F5F8" if i % 2 else "FFFFFF"))
            props.append(fill)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                value = OxmlElement(f"w:{edge}")
                value.set(qn("w:w"), "85")
                value.set(qn("w:type"), "dxa")
                margins.append(value)
            props.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.03
                if j == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(10)
                    run.font.bold = i == 0 or j == 0
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else "183345")
    anchor._p.addprevious(table._tbl)


def main():
    document = Document(SOURCE)
    original = list(document.paragraphs)
    anchor = next(p for p in original if p.text.startswith("2.6 Aplicación"))
    # Renumerar las subsecciones posteriores y el gráfico de casos existente.
    for paragraph in original:
        for old, new in (("2.9 ", "2.10 "), ("2.8 ", "2.9 "), ("2.7 ", "2.8 "), ("2.6 ", "2.7 ")):
            if paragraph.text.startswith(old):
                paragraph.runs[0].text = paragraph.runs[0].text.replace(old, new, 1)
                break
        if paragraph.text.startswith("Figura 2.2."):
            paragraph.runs[0].text = paragraph.runs[0].text.replace("Figura 2.2.", "Figura 2.3.", 1)
    evidence = next(p for p in original if p.text.startswith("Evidencia: diagrama lógico"))
    evidence.text = "Evidencia: diagramas de componentes y eventos, interfaces, matriz de responsabilidades y casos compartidos con el recorrido de Benjamín."

    paragraph_before(anchor, "2.6 Eventos que articulan la arquitectura", "Heading 3", True)
    paragraph_before(anchor,
        "Los eventos muestran qué hecho activa la colaboración entre módulos. La aceptación de la cotización permite que Ventas solicite la evaluación; su recepción activa el control de transición. C04 conserva una respuesta motivada y C07 registra su trazabilidad. Los identificadores E01 a E07 sirven para explicar el diseño y no agregan operaciones al contrato 1.0.")
    image_paragraph = paragraph_before(anchor, "")
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_with_next = True
    shape = image_paragraph.add_run().add_picture(str(IMAGE), width=Inches(6.98))
    shape._inline.docPr.set("descr", "Eventos E01 a E07, responsables y respuestas de CRM ERP. La señal conceptual REP-EP-001 conecta O2C con P2P sin compra automática.")
    caption = paragraph_before(anchor,
        "Figura 2.2. Eventos, módulos y respuestas del caso presencial. Las ramas son variantes independientes. Elaboración propia asistida por IA, basada en las evidencias locales EP-N01, EP-S01 y EP-E01.", size=9)
    caption.paragraph_format.space_after = Pt(4)

    paragraph_before(anchor, "Registro y respuesta de los eventos", "Heading 3", True)
    paragraph_before(anchor,
        "La matriz relaciona cada hecho con su origen, quién lo atiende y la evidencia disponible. Un evento es un hecho ocurrido; I03 es una interfaz y ACEPTADA es un estado de respuesta. La aceptación del cliente y la validación de la solicitud son hechos distintos.")
    events_table(document, anchor)
    paragraph_before(anchor, "")
    paragraph_before(anchor,
        "La bitácora técnica conserva event_id, company_id, request_id, correlation_id, actor_id, interface_id, event_type, code y occurred_at. SOL identifica la solicitud y COR reúne los hechos del recorrido. La decisión y su bitácora se guardan en una misma transacción según el diseño.", size=10)
    paragraph_before(anchor,
        "E01 pertenece a la narrativa sintética previa. E02 a E05 y E07 se contrastan con la bitácora del simulador. E06 es una señal conceptual documentada por el arnés tras la respuesta parcial: no crea compra, reserva ni fecha automática. La vista no presupone un bus de mensajes instalado.", size=10)
    paragraph_before(anchor,
        "El diagrama desarrolla las tres variantes del ejemplo. PENDIENTE por dependencia inaccesible y su recuperación permanecen definidos en el contrato. La correspondencia literal con la rúbrica oficial requiere revisar el documento docente original.", size=10)

    for paragraph in original:
        if paragraph.text.startswith("02:15 a 03:00"):
            paragraph.text = "02:15 a 03:00  Eventos y controles del caso normal. “La aceptación de COT-DEMO-001 v1 permite solicitar la evaluación. Al recibirla, C04 verifica oferta, autorización y stock. Cinco autoradios a 400 soles, con 5 % de descuento, suman 1,900. El evento de decisión queda correlacionado en la bitácora: estado ACEPTADA y código VALIDADA, con order_id nulo. Todavía no confirma pedido.”"
        if paragraph.text.startswith("03:00 a 03:35"):
            paragraph.text = "03:00 a 03:35  Evento que conecta con abastecimiento. “Si hay tres unidades para cinco solicitadas, la evaluación devuelve disponibilidad parcial. Después, el arnés documenta REP-EP-001 como necesidad trazable para que dueño y abastecimiento evalúen alternativas. Este evento conceptual conecta O2C con P2P; no crea una compra ni permite prometer fecha.”"
        if paragraph.text.startswith("Contribución verificable. Diagrama lógico"):
            paragraph.text = "Contribución verificable. Diagramas de componentes C01 a C08 y eventos E01 a E07, interfaces I01 a I07, matriz de sistemas de registro, continuidad de PC2 y correspondencia con los casos de Benjamín. Decisión defendida: mantener Odoo integrado con responsabilidades, eventos y controles explícitos antes de comprometer la operación."

    # Conservar el resto de la plantilla y la sección completa de Benjamín.
    before_boundary = next(i for i, p in enumerate(original) if p.text.startswith("3. Recorrido"))
    expected_tail = [p._p.xml for p in original[before_boundary:]]
    actual_paragraphs = list(document.paragraphs)
    after_boundary = next(i for i, p in enumerate(actual_paragraphs) if p.text.startswith("3. Recorrido"))
    assert expected_tail == [p._p.xml for p in actual_paragraphs[after_boundary:]]
    assert len(document.inline_shapes) == 9
    assert len(document.tables) == 18
    document.save(DEST)
    print(f"Word actualizado: {DEST}")
    print("9 figuras; 18 tablas; sección 3 y plantilla posterior conservadas.")


if __name__ == "__main__":
    main()
